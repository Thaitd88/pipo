"""
Script nhập và cập nhật menu hàng loạt từ file Excel Menu_PiPo_Garden_Cap_Nhat.xlsx
Chạy lệnh: python scripts/import_menu_from_excel.py [file_path]
"""

import sys, os, sqlite3, json, re
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

excel_file = sys.argv[1] if len(sys.argv) > 1 else "Menu_PiPo_Garden_Cap_Nhat.xlsx"

if not os.path.exists(excel_file):
    print(f"❌ Không tìm thấy file: {excel_file}")
    sys.exit(1)

wb = openpyxl.load_workbook(excel_file, data_only=True)
if "Thực Đơn PiPo Garden" not in wb.sheetnames:
    print(f"❌ File Excel không có sheet 'Thực Đơn PiPo Garden'")
    sys.exit(1)

ws = wb["Thực Đơn PiPo Garden"]

conn = sqlite3.connect("data/pipo_garden.db")
conn.row_factory = sqlite3.Row
cur = conn.cursor()

# Get existing categories map
cur.execute("SELECT id, name FROM categories")
cat_map = {}
for r in cur.fetchall():
    cat_map[r["name"].strip()] = r["id"]
    # strip icons
    clean_cat = re.sub(r'[^\w\s-]', '', r["name"]).strip()
    if clean_cat:
        cat_map[clean_cat] = r["id"]

# Station reverse map
station_rev = {
    "Bếp Nóng": "HOT_KITCHEN",
    "Bếp Lạnh": "COLD_KITCHEN",
    "Quầy Bar": "BAR",
    "Sân Pickleball": "PICKLE",
    "HOT_KITCHEN": "HOT_KITCHEN",
    "COLD_KITCHEN": "COLD_KITCHEN",
    "BAR": "BAR",
    "PICKLE": "PICKLE"
}

updated_count = 0
inserted_count = 0

for r in range(4, ws.max_row + 1):
    dish_id_val = ws.cell(r, 1).value
    cat_name = str(ws.cell(r, 2).value or "").strip()
    dish_name = str(ws.cell(r, 3).value or "").strip()
    taste_val = str(ws.cell(r, 4).value or "").strip()
    portion = str(ws.cell(r, 5).value or "").strip()
    price_val = ws.cell(r, 7).value
    station_val = str(ws.cell(r, 10).value or "").strip()
    menu_type_val = str(ws.cell(r, 11).value or "").strip()
    status_val = str(ws.cell(r, 12).value or "").strip()
    img_file = str(ws.cell(r, 13).value or "").strip()
    desc_val = str(ws.cell(r, 16).value or "").strip()

    if not dish_name:
        continue

    try:
        price = int(float(price_val or 0))
    except Exception:
        price = 0

    station = station_rev.get(station_val, "HOT_KITCHEN")
    is_avail = 0 if "ngưng" in status_val.lower() or "tạm" in status_val.lower() or status_val == "0" else 1

    # Image URL
    image_url = ""
    if img_file and not img_file.startswith("(") and img_file != "Chưa có ảnh":
        if img_file == "logo.png":
            image_url = "/static/logo.png"
        else:
            image_url = f"/static/dishes/{img_file}"

    # Parse options_json
    options_json = "[]"
    if taste_val and taste_val.lower() != "tiêu chuẩn":
        # Check if format is "Chọn vị chế biến: X, Y, Z"
        if ":" in taste_val:
            parts = taste_val.split(":", 1)
            title = parts[0].strip()
            opts = [o.strip() for o in parts[1].split(",") if o.strip()]
        else:
            title = "Chọn vị chế biến"
            opts = [o.strip() for o in taste_val.split(",") if o.strip()]
        if opts:
            options_json = json.dumps([{"name": title, "options": opts}], ensure_ascii=False)

    # Find category_id
    cat_id = 1
    clean_cat_input = re.sub(r'[^\w\s-]', '', cat_name).strip()
    for k, v in cat_map.items():
        if k in cat_name or k in clean_cat_input or clean_cat_input in k:
            cat_id = v
            break

    # Determine if update or insert
    is_numeric_id = False
    try:
        dish_id = int(dish_id_val)
        is_numeric_id = True
    except (ValueError, TypeError):
        dish_id = None

    existing_dish = None
    if is_numeric_id:
        cur.execute("SELECT id FROM menu_items WHERE id = ?", (dish_id,))
        existing_dish = cur.fetchone()
    if not existing_dish:
        cur.execute("SELECT id FROM menu_items WHERE LOWER(name) = LOWER(?)", (dish_name,))
        existing_dish = cur.fetchone()

    if existing_dish:
        cur_id = existing_dish["id"]
        updates = ["name = ?", "price = ?", "portion = ?", "options_json = ?", "station = ?", "is_available = ?"]
        params = [dish_name, price, portion, options_json, station, is_avail]
        if image_url:
            updates.append("image_url = ?")
            params.append(image_url)
        if desc_val:
            updates.append("description = ?")
            params.append(desc_val)
        if cat_id:
            updates.append("category_id = ?")
            params.append(cat_id)
        
        params.append(cur_id)
        cur.execute(f"UPDATE menu_items SET {', '.join(updates)} WHERE id = ?", tuple(params))
        updated_count += 1
    else:
        cur.execute("""
            INSERT INTO menu_items (
                category_id, name, price, portion, options_json, station,
                is_available, image_url, description, main_menu_type
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            cat_id, dish_name, price, portion, options_json, station,
            is_avail, image_url or "/static/logo.png", desc_val, "RESTAURANT"
        ))
        inserted_count += 1

conn.commit()
conn.close()

print(f"✅ ĐÃ NHẬP MENU TỪ EXCEL THÀNH CÔNG!")
print(f"  • Món đã cập nhật: {updated_count}")
print(f"  • Món mới tạo thêm: {inserted_count}")
