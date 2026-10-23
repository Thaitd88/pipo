import sys, os, sqlite3, json, re
import urllib.parse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

# Database connection
db_path = "data/pipo_garden.db"
conn = sqlite3.connect(db_path)
conn.row_factory = sqlite3.Row
cur = conn.cursor()

wb = Workbook()

# -------------------------------------------------------------------------
# STYLES DEFINITION
# -------------------------------------------------------------------------
font_title = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
fill_title = PatternFill(start_color="1B4D3E", end_color="1B4D3E", fill_type="solid") # Dark Forest Green

font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color="235D4B", end_color="235D4B", fill_type="solid")

font_header_sec = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
fill_header_sec = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Navy

font_data = Font(name="Segoe UI", size=10, color="000000")
font_data_bold = Font(name="Segoe UI", size=10, bold=True, color="000000")
font_id = Font(name="Consolas", size=10, bold=True, color="1E3A8A")
font_link = Font(name="Segoe UI", size=9, color="0066CC", underline="single")
font_notice = Font(name="Segoe UI", size=9, italic=True, color="666666")

fill_zebra = PatternFill(start_color="F4F8F6", end_color="F4F8F6", fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
fill_highlight = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Light Amber

thin_border_side = Side(border_style="thin", color="D1D5DB")
border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
border_header = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_right = Alignment(horizontal="right", vertical="center")

# Base project dir
base_dir = os.path.abspath(".")
static_dishes_dir = os.path.join(base_dir, "app", "static", "dishes")
static_dir = os.path.join(base_dir, "app", "static")

# -------------------------------------------------------------------------
# SHEET 1: MENU THỰC ĐƠN PIPO GARDEN
# -------------------------------------------------------------------------
ws1 = wb.active
ws1.title = "Thực Đơn PiPo Garden"
ws1.views.sheetView[0].showGridLines = True

# Title row
ws1.merge_cells("A1:P1")
title_cell = ws1["A1"]
title_cell.value = "BẢNG QUẢN LÝ & CẬP NHẬT THỰC ĐƠN PIPO GARDEN PHÚ QUỐC (F&B & SÂN PICKLEBALL)"
title_cell.font = font_title
title_cell.fill = fill_title
title_cell.alignment = Alignment(horizontal="center", vertical="center")
ws1.row_dimensions[1].height = 40

# Subtitle / Info row
ws1.merge_cells("A2:P2")
sub_cell = ws1["A2"]
sub_cell.value = "Quý khách / Chủ quán chỉnh sửa trực tiếp các cột Tên món, Vị, Định lượng, Giá bán, Giá cost, Tên ảnh... sau đó lưu file lại để hệ thống cập nhật đồng loạt."
sub_cell.font = font_notice
sub_cell.fill = fill_zebra
sub_cell.alignment = Alignment(horizontal="center", vertical="center")
ws1.row_dimensions[2].height = 24

# Header row 3
headers1 = [
    ("Mã món (ID)", 12, "CENTER"),
    ("Nhóm món", 18, "LEFT"),
    ("Tên món ăn / Dịch vụ", 28, "LEFT"),
    ("Vị / Cách chế biến", 32, "LEFT"),
    ("Quy cách / Định lượng phần", 26, "LEFT"),
    ("Định lượng NVL chính (BOM)", 36, "LEFT"),
    ("Giá bán (VNĐ)", 16, "RIGHT"),
    ("Giá cost NVL (VNĐ)", 16, "RIGHT"),
    ("% Cost/Giá", 12, "CENTER"),
    ("Quầy chế biến", 16, "CENTER"),
    ("Phân hệ Menu", 18, "CENTER"),
    ("Trạng thái bán", 14, "CENTER"),
    ("Tên file ảnh trong folder", 30, "LEFT"),
    ("Đường dẫn ảnh cục bộ trên máy", 36, "LEFT"),
    ("Link ảnh Web (Online)", 34, "LEFT"),
    ("Mô tả món ăn", 35, "LEFT")
]

ws1.row_dimensions[3].height = 30
for col_idx, (h_name, width, _) in enumerate(headers1, 1):
    cell = ws1.cell(row=3, column=col_idx, value=h_name)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = border_header
    col_letter = get_column_letter(col_idx)
    ws1.column_dimensions[col_letter].width = width

# Fetch dishes
cur.execute('''
    SELECT m.id, m.category_id, c.name as category_name, c.icon as category_icon,
           m.name, m.price, m.portion, m.station, m.is_available, m.options_json,
           m.image_url, m.main_menu_type, m.description
    FROM menu_items m
    LEFT JOIN categories c ON m.category_id = c.id
    ORDER BY m.category_id ASC, m.id ASC
''')
menu_rows = cur.fetchall()

station_map = {
    "HOT_KITCHEN": "Bếp Nóng",
    "COLD_KITCHEN": "Bếp Lạnh",
    "BAR": "Quầy Bar",
    "PICKLE": "Sân Pickleball"
}

main_menu_map = {
    "RESTAURANT": "Nhà Hàng",
    "LUNCH": "Cơm Trưa",
    "DRINK_BREAKFAST": "Ăn Sáng - Cafe",
    "PICKLE": "Sân Pickleball",
    "RESTAURANT,PICKLE": "Nhà Hàng & Sân Pick"
}

row_idx = 4
for r in menu_rows:
    dish_id = r["id"]
    cat_name = f"{r['category_icon'] or ''} {r['category_name'] or ''}".strip()
    dish_name = r["name"] or ""
    price = r["price"] or 0
    portion = r["portion"] or ""
    desc = r["description"] or ""
    is_avail = "Đang bán" if r["is_available"] == 1 else "Tạm ngưng"
    station_txt = station_map.get(r["station"], r["station"] or "")
    
    # Menu types
    m_types = [main_menu_map.get(t.strip(), t.strip()) for t in (r["main_menu_type"] or "RESTAURANT").split(",")]
    menu_type_txt = ", ".join(m_types)

    # Parse options / Taste
    opt_str = ""
    try:
        raw_opt = json.loads(r["options_json"] or "[]")
        if isinstance(raw_opt, list) and raw_opt:
            parts = []
            for item in raw_opt:
                if isinstance(item, dict):
                    title = item.get("name", "")
                    sub = item.get("options", []) or item.get("values", [])
                    if sub:
                        parts.append(f"{title}: {', '.join(sub)}")
                    elif title:
                        parts.append(title)
                elif isinstance(item, str):
                    parts.append(item)
            opt_str = "; ".join(parts)
    except Exception:
        opt_str = str(r["options_json"] or "")
    if not opt_str:
        opt_str = "Tiêu chuẩn"

    # Fetch BOM
    cur.execute('''
        SELECT ii.name as ing_name, rb.quantity, ii.unit, ii.cost_price, (rb.quantity * ii.cost_price) as item_cost
        FROM recipe_bom rb
        JOIN inventory_items ii ON rb.ingredient_id = ii.id
        WHERE rb.menu_item_id = ?
    ''', (dish_id,))
    bom_rows = cur.fetchall()
    if bom_rows:
        bom_desc = "; ".join([f"{b['ing_name']}: {b['quantity']} {b['unit']}" for b in bom_rows])
        total_cost = sum(b["item_cost"] or 0 for b in bom_rows)
    else:
        bom_desc = "Chưa thiết lập BOM"
        total_cost = 0

    # Image info
    img_url_raw = r["image_url"] or ""
    # Strip ?v=...
    clean_url = img_url_raw.split("?")[0].strip()
    
    img_filename = ""
    local_path = ""
    online_url = ""

    if clean_url:
        img_filename = os.path.basename(clean_url)
        if clean_url.startswith("/static/dishes/"):
            local_path = os.path.join(static_dishes_dir, img_filename)
            online_url = f"https://pipogarden.com/static/dishes/{img_filename}"
        elif clean_url.startswith("/static/"):
            local_path = os.path.join(static_dir, img_filename)
            online_url = f"https://pipogarden.com/static/{img_filename}"
        elif clean_url.startswith("http"):
            online_url = clean_url
            img_filename = os.path.basename(urllib.parse.urlparse(clean_url).path)
            local_path = "(Ảnh link online ngoài)"
        else:
            local_path = os.path.join(static_dir, clean_url.lstrip("/"))
            online_url = f"https://pipogarden.com/{clean_url.lstrip('/')}"
    else:
        img_filename = "(Chưa có ảnh)"
        local_path = "(Chưa có ảnh)"
        online_url = ""

    fill_curr = fill_zebra if row_idx % 2 == 0 else fill_white

    row_data = [
        (dish_id, align_center, font_id, None),
        (cat_name, align_left, font_data_bold, None),
        (dish_name, align_left, font_data_bold, None),
        (opt_str, align_left, font_data, None),
        (portion, align_left, font_data, None),
        (bom_desc, align_left, font_data, None),
        (price, align_right, font_data_bold, '#,##0 "đ"'),
        (int(total_cost), align_right, font_data, '#,##0 "đ"'),
        (f"=IF(G{row_idx}>0, H{row_idx}/G{row_idx}, 0)", align_center, font_data, '0.0%'),
        (station_txt, align_center, font_data, None),
        (menu_type_txt, align_center, font_data, None),
        (is_avail, align_center, font_data, None),
        (img_filename, align_left, font_data, None),
        (local_path, align_left, font_link if local_path.startswith("d:") or local_path.startswith("D:") else font_data, None),
        (online_url, align_left, font_link if online_url.startswith("http") else font_data, None),
        (desc, align_left, font_data, None),
    ]

    ws1.row_dimensions[row_idx].height = 24

    for col_i, (val, alignment, font_style, num_format) in enumerate(row_data, 1):
        c = ws1.cell(row=row_idx, column=col_i)
        if str(val).startswith("="):
            c.value = val
        else:
            c.value = val
        c.font = font_style
        c.alignment = alignment
        c.fill = fill_curr
        c.border = border_cell
        if num_format:
            c.number_format = num_format
        
        # Hyperlink for online URL
        if col_i == 15 and online_url.startswith("http"):
            c.hyperlink = online_url
        # Hyperlink for local file if exists
        if col_i == 14 and os.path.exists(local_path):
            c.hyperlink = f"file:///{local_path.replace(os.sep, '/')}"

    row_idx += 1

# Auto-filter on Sheet 1
ws1.auto_filter.ref = f"A3:P{row_idx - 1}"
ws1.freeze_panes = "A4"


# -------------------------------------------------------------------------
# SHEET 2: NGUYÊN VẬT LIỆU KHO & GIÁ VỐN (BOM INVENTORY)
# -------------------------------------------------------------------------
ws2 = wb.create_sheet(title="Nguyên Vật Liệu Kho")
ws2.views.sheetView[0].showGridLines = True

# Title row
ws2.merge_cells("A1:G1")
title_cell2 = ws2["A1"]
title_cell2.value = "DANH MỤC NGUYÊN VẬT LIỆU KHO & GIÁ VỐN COST (PIPO GARDEN)"
title_cell2.font = font_title
title_cell2.fill = fill_header_sec
title_cell2.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[1].height = 36

headers2 = [
    ("Mã NVL", 14, "CENTER"),
    ("Tên Nguyên Vật Liệu", 34, "LEFT"),
    ("Đơn vị tính", 14, "CENTER"),
    ("Giá vốn cost (VNĐ)", 18, "RIGHT"),
    ("Tồn kho hiện tại", 16, "RIGHT"),
    ("Tồn tối thiểu", 16, "RIGHT"),
    ("Nhà cung cấp / Ghi chú", 30, "LEFT")
]

ws2.row_dimensions[2].height = 28
for col_idx, (h_name, width, _) in enumerate(headers2, 1):
    cell = ws2.cell(row=2, column=col_idx, value=h_name)
    cell.font = font_header
    cell.fill = fill_header_sec
    cell.alignment = align_center
    cell.border = border_header
    ws2.column_dimensions[get_column_letter(col_idx)].width = width

cur.execute("SELECT id, code, name, unit, cost_price, stock_quantity, min_stock, supplier FROM inventory_items ORDER BY id ASC")
inv_rows = cur.fetchall()

row_idx2 = 3
for inv in inv_rows:
    fill_curr = fill_zebra if row_idx2 % 2 == 0 else fill_white
    ws2.row_dimensions[row_idx2].height = 22

    row_data2 = [
        (inv["code"] or f"NVL{inv['id']:02d}", align_center, font_id, None),
        (inv["name"] or "", align_left, font_data_bold, None),
        (inv["unit"] or "", align_center, font_data, None),
        (inv["cost_price"] or 0, align_right, font_data_bold, '#,##0 "đ"'),
        (inv["stock_quantity"] or 0, align_right, font_data, '#,##0.00'),
        (inv["min_stock"] or 0, align_right, font_data, '#,##0.00'),
        (inv["supplier"] or "Chợ Dương Đông / Phú Quốc", align_left, font_data, None)
    ]

    for col_i, (val, alignment, font_style, num_format) in enumerate(row_data2, 1):
        c = ws2.cell(row=row_idx2, column=col_i, value=val)
        c.font = font_style
        c.alignment = alignment
        c.fill = fill_curr
        c.border = border_cell
        if num_format:
            c.number_format = num_format

    row_idx2 += 1

ws2.auto_filter.ref = f"A2:G{row_idx2 - 1}"
ws2.freeze_panes = "A3"


# -------------------------------------------------------------------------
# SHEET 3: HƯỚNG DẪN CẬP NHẬT (GUIDE)
# -------------------------------------------------------------------------
ws3 = wb.create_sheet(title="Hướng Dẫn Cập Nhật")
ws3.views.sheetView[0].showGridLines = True

ws3.merge_cells("A1:E1")
title_cell3 = ws3["A1"]
title_cell3.value = "HƯỚNG DẪN CHỈNH SỬA & QUY TRÌNH CẬP NHẬT MENU HÀNG LOẠT"
title_cell3.font = font_title
title_cell3.fill = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
title_cell3.alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[1].height = 36

guide_items = [
    ("1. Cột 'Mã món (ID)'", "Tuyệt đối giữ nguyên mã số này đối với các món đã có. Hệ thống sẽ căn cứ vào ID này để cập nhật giá, vị, ảnh chính xác 100% vào cơ sở dữ liệu."),
    ("2. Cột 'Nhóm món'", "Có thể chọn 1 trong các nhóm chuẩn: Khai Vị, Gà - Vịt, Heo - Bò, Hải Sản, Cơm - Mì, Đồ Uống, Hot BBQ, Dịch Vụ Sân Pickleball."),
    ("3. Cột 'Tên món ăn'", "Gõ tên hiển thị mong muốn trên menu POS và Menu QR của khách hàng."),
    ("4. Cột 'Vị / Cách chế biến'", "Ghi các tùy chọn cách chế biến khách có thể chọn khi gọi món. Cách nhau bởi dấu phẩy.\nVí dụ: Nướng muối ớt, Hấp mắm nhĩ, Chiên mắm, Lẩu thuốc bắc.\nNếu món không có vị tùy chọn, ghi 'Tiêu chuẩn'."),
    ("5. Cột 'Quy cách / Định lượng'", "Mô tả phần ăn hiển thị cho khách và bếp (Ví dụ: Gà 1 con 1.3kg, Đĩa 300g, 1 lon 320ml, 1 giờ...)."),
    ("6. Cột 'Giá bán (VNĐ)'", "Nhập số tiền bán thực tế (ví dụ: 99000, 350000). Không cần nhập chữ 'đ'."),
    ("7. Cột 'Quầy chế biến'", "Nhập một trong các quầy sau để máy in chia đúng trạm:\n- Bếp Nóng (Món xào, nướng, lẩu, chiên -> In máy bếp 251)\n- Bếp Lạnh (Món gỏi, salad -> In quầy thu ngân 252)\n- Quầy Bar (Bia, nước ngọt, trà, nước ép -> In quầy thu ngân 252)\n- Sân Pickleball (Dịch vụ sân, thuê vợt, bóng -> Quầy sân pick 252)"),
    ("8. Cột 'Tên file ảnh trong folder'", "Tên file ảnh tương ứng đặt trong thư mục: d:\\AI\\Restaurant\\app\\static\\dishes\\\nVí dụ: ga_nuong_muoi_ot.jpg, bo_xuyen_tieu.jpg, nuoc_suoi.jpg.\nBạn chỉ cần copy ảnh vào thư mục đó và ghi tên file ảnh vào cột này."),
    ("9. Thêm món mới vào menu", "Nếu muốn thêm món mới: Thêm dòng mới ở cuối bảng, cột 'Mã món (ID)' để chữ 'MỚI' hoặc để trống, điền đủ Tên món, Nhóm, Giá, Ảnh..."),
    ("10. Sau khi chỉnh sửa xong", "Lưu file Excel lại và thông báo cho Antigravity. Hệ thống sẽ tự động đọc file và nạp toàn bộ menu mới vào CSDL máy Thu ngân & tự động Deploy lên Cloud pipogarden.com ngay lập tức!")
]

ws3.column_dimensions["A"].width = 28
ws3.column_dimensions["B"].width = 90

h3_cell1 = ws3.cell(row=2, column=1, value="Hạng mục / Cột dữ liệu")
h3_cell1.font = font_header
h3_cell1.fill = PatternFill(start_color="4B5563", end_color="4B5563", fill_type="solid")
h3_cell1.alignment = align_center

h3_cell2 = ws3.cell(row=2, column=2, value="Hướng dẫn chi tiết & Lưu ý quan trọng")
h3_cell2.font = font_header
h3_cell2.fill = PatternFill(start_color="4B5563", end_color="4B5563", fill_type="solid")
h3_cell2.alignment = align_center
ws3.row_dimensions[2].height = 28

for idx, (col_name, desc_text) in enumerate(guide_items, 3):
    ws3.row_dimensions[idx].height = 42 if "\n" in desc_text else 28
    fill_g = fill_zebra if idx % 2 == 0 else fill_white

    c1 = ws3.cell(row=idx, column=1, value=col_name)
    c1.font = font_data_bold
    c1.fill = fill_g
    c1.alignment = align_left
    c1.border = border_cell

    c2 = ws3.cell(row=idx, column=2, value=desc_text)
    c2.font = font_data
    c2.fill = fill_g
    c2.alignment = align_left
    c2.border = border_cell

# Save workbook
output_path = os.path.abspath("Menu_PiPo_Garden_Cap_Nhat.xlsx")
wb.save(output_path)
print(f"EXCEL CREATED SUCCESSFULLY AT: {output_path}")

conn.close()
